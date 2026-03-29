"""Diagnose HDBSCAN clustering for a specific fixture file.

Run from repo root:
    .venv/bin/python scripts/diagnose_hdbscan.py
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import numpy as np
from openpyxl import load_workbook

from api.python_analyze import (
    DOMINANT_CLUSTER_REASSIGNMENT_SIMILARITY_THRESHOLD,
    HDBSCAN_RELAXED_DOMINANT_SIMILARITY_THRESHOLD,
    HDBSCAN_RELAXED_MAX_DOMINANT_CLUSTER_SHARE,
    HDBSCAN_RELAXED_MAX_NOISE_RATIO,
    HDBSCAN_RELAXED_MIN_CLUSTER_COUNT,
    HDBSCAN_RELAXED_NOISE_MARGIN_THRESHOLD,
    HDBSCAN_RELAXED_NOISE_SIMILARITY_THRESHOLD,
    HDBSCAN_SEARCH_MAX_CLUSTER_COUNT_CAP,
    HDBSCAN_SEARCH_MAX_CLUSTER_COUNT_FACTOR,
    HDBSCAN_SEARCH_MAX_DOMINANT_CLUSTER_SHARE,
    HDBSCAN_SEARCH_MAX_NOISE_RATIO,
    HDBSCAN_SEARCH_MIN_AVG_CLUSTER_SIZE_FACTOR,
    HDBSCAN_SEARCH_MIN_CLUSTER_COUNT,
    NOISE_REASSIGNMENT_MARGIN_THRESHOLD,
    NOISE_REASSIGNMENT_SIMILARITY_THRESHOLD,
    build_embeddings,
    build_hdbscan_parameter_grid,
    clean_text,
    cluster_with_hdbscan,
    evaluate_hdbscan_candidate,
    is_tag_like,
    reassign_noise_points,
    reduce_dimensions,
)
from openai import OpenAI
import os

FIXTURE = Path("fixtures/analysis-inputs/Verbatim IT extraction 2025-2026.xlsx")


def load_texts() -> list[str]:
    wb = load_workbook(str(FIXTURE), read_only=True, data_only=True)
    ws = wb["Sheet2"]
    headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    text_col_idx = None
    for idx, h in enumerate(headers):
        if h.lower() == "verbatim_comment":
            text_col_idx = idx
            break
    if text_col_idx is None:
        raise RuntimeError(f"Column 'verbatim_comment' not found in {headers}")

    texts: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw = str(row[text_col_idx] or "").strip()
        cleaned = clean_text(raw)
        if not cleaned:
            continue
        if is_tag_like(cleaned):
            continue
        texts.append(cleaned)
    wb.close()
    return texts


def print_grid(
    label: str,
    grid: list[tuple[int, int]],
    reduced: np.ndarray,
    embeddings: np.ndarray,
    *,
    sim_thresh: float,
    margin_thresh: float,
    dom_sim_thresh: float,
    max_noise: float,
    max_dominant: float,
    min_cc: int | None,
) -> int:
    n = len(reduced)
    sqrt_n = math.sqrt(n)
    effective_min_cc = min_cc if min_cc is not None else (HDBSCAN_SEARCH_MIN_CLUSTER_COUNT if n >= 15 else 2)
    max_cc = max(effective_min_cc + 1, min(HDBSCAN_SEARCH_MAX_CLUSTER_COUNT_CAP, math.ceil(sqrt_n * HDBSCAN_SEARCH_MAX_CLUSTER_COUNT_FACTOR)))
    min_avg = max(5.0, sqrt_n * HDBSCAN_SEARCH_MIN_AVG_CLUSTER_SIZE_FACTOR)

    print(f"\n{'='*80}")
    print(f"  {label}")
    print(f"{'='*80}")
    print(f"  cluster_count: [{effective_min_cc}, {max_cc}]  |  max noise: {max_noise}  |  max dominant: {max_dominant}")
    print(f"  noise reassign sim: {sim_thresh}  |  margin: {margin_thresh}  |  dominant sim: {dom_sim_thresh}")
    print(f"  min avg cluster size: {min_avg:.2f}")
    print()
    print(f"{'mcs':>5} {'ms':>4} │ {'clusters':>8} {'noise%':>7} {'dominant%':>9} {'avg_size':>8} {'conf':>6} │ {'result'}")
    print("─" * 90)

    valid_count = 0
    for mcs, ms in grid:
        labels, conf = cluster_with_hdbscan(reduced, mcs, ms)
        labels_copy = labels.copy()
        conf_copy = conf.copy()
        labels_copy, conf_copy, _ = reassign_noise_points(
            embeddings, labels_copy, conf_copy,
            similarity_threshold=sim_thresh,
            margin_threshold=margin_thresh,
            dominant_similarity_threshold=dom_sim_thresh,
        )
        ev = evaluate_hdbscan_candidate(
            labels_copy, conf_copy,
            max_noise_ratio=max_noise,
            max_dominant_share=max_dominant,
            min_cluster_count_override=min_cc,
        )

        fails = []
        if ev["cluster_count"] < effective_min_cc:
            fails.append(f"clusters<{effective_min_cc}")
        if ev["cluster_count"] > max_cc:
            fails.append(f"clusters>{max_cc}")
        if ev["avg_cluster_size"] < min_avg:
            fails.append(f"avg<{min_avg:.0f}")
        if ev["noise_ratio"] > max_noise:
            fails.append(f"noise>{max_noise}")
        if ev["dominant_cluster_share"] > max_dominant:
            fails.append(f"dom>{max_dominant}")

        status = "VALID" if not fails else "FAIL: " + ", ".join(fails)
        if not fails:
            valid_count += 1

        print(
            f"{mcs:>5} {ms:>4} │ "
            f"{ev['cluster_count']:>8} "
            f"{ev['noise_ratio']:>6.1%} "
            f"{ev['dominant_cluster_share']:>8.1%} "
            f"{ev['avg_cluster_size']:>8.1f} "
            f"{ev['clustered_confidence']:>6.3f} │ "
            f"{status}"
        )

    print("─" * 90)
    print(f"  Valid candidates: {valid_count}/{len(grid)}")
    return valid_count


def main() -> None:
    print(f"Loading texts from {FIXTURE} ...")
    texts = load_texts()
    n = len(texts)
    print(f"Clusterable comments: {n}")
    print(f"sqrt(n) = {math.sqrt(n):.2f}")

    api_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not api_key:
        print("ERROR: OPENAI_API_KEY not set.")
        sys.exit(1)

    client = OpenAI(api_key=api_key)
    print(f"\nGenerating embeddings for {n} texts ...")
    embeddings = build_embeddings(client, texts)
    print(f"Embeddings shape: {embeddings.shape}")

    reduced = reduce_dimensions(embeddings)
    print(f"Reduced shape: {reduced.shape}")

    grid = build_hdbscan_parameter_grid(n)
    print(f"Parameter grid: {len(grid)} combinations")

    strict_valid = print_grid(
        "STRICT PASS (original thresholds, normal noise reassignment)",
        grid, reduced, embeddings,
        sim_thresh=NOISE_REASSIGNMENT_SIMILARITY_THRESHOLD,
        margin_thresh=NOISE_REASSIGNMENT_MARGIN_THRESHOLD,
        dom_sim_thresh=DOMINANT_CLUSTER_REASSIGNMENT_SIMILARITY_THRESHOLD,
        max_noise=HDBSCAN_SEARCH_MAX_NOISE_RATIO,
        max_dominant=HDBSCAN_SEARCH_MAX_DOMINANT_CLUSTER_SHARE,
        min_cc=None,
    )

    relaxed_valid = print_grid(
        "RELAXED PASS (raised noise tolerance, aggressive reassignment)",
        grid, reduced, embeddings,
        sim_thresh=HDBSCAN_RELAXED_NOISE_SIMILARITY_THRESHOLD,
        margin_thresh=HDBSCAN_RELAXED_NOISE_MARGIN_THRESHOLD,
        dom_sim_thresh=HDBSCAN_RELAXED_DOMINANT_SIMILARITY_THRESHOLD,
        max_noise=HDBSCAN_RELAXED_MAX_NOISE_RATIO,
        max_dominant=HDBSCAN_RELAXED_MAX_DOMINANT_CLUSTER_SHARE,
        min_cc=HDBSCAN_RELAXED_MIN_CLUSTER_COUNT,
    )

    print(f"\n{'='*80}")
    if strict_valid > 0:
        print(f"  RESULT: Strict pass has {strict_valid} valid candidate(s) — HDBSCAN would be used.")
    elif relaxed_valid > 0:
        print(f"  RESULT: Relaxed pass has {relaxed_valid} valid candidate(s) — HDBSCAN would be used (relaxed).")
    else:
        print("  RESULT: Both passes failed — KMeans fallback would be used.")
    print(f"{'='*80}")


if __name__ == "__main__":
    main()
